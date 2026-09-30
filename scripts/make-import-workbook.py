#!/usr/bin/env python3
"""Builds the clean-rebuild import workbook for the HR OS Google Sheet.

usage: python3 scripts/make-import-workbook.py <old-workbook.xlsx> <out.xlsx>

Reads the old VFL_HR_OS workbook (read-only) and writes ONLY the tabs the new Apps Script needs. The output contains sensitive
statutory IDs (EMPLOYEE_STATUTORY_IDS): it must NEVER be committed (*.xlsx is in .gitignore). Only counts are printed.
Import it in Google Sheets with File > Import > Upload > "Insert new sheet(s)".
"""
import json
import re
import sys

import openpyxl

# SCHEMA-BEGIN (kept in sync with apps-script by tests/import-workbook.test.js; do not edit by hand)
SCHEMA = json.loads(r"""{
 "headers": {
  "PAYROLL_CONTROL": [
   "KEY",
   "VALUE",
   "NOTE",
   "UPDATED_AT"
  ],
  "PAYROLL_PERIOD_CATEGORY": [
   "PAYROLL_MONTH",
   "PAYROLL_CATEGORY",
   "WORKING_DAYS",
   "STATUS",
   "APPROVED_BY",
   "APPROVED_AT",
   "NOTE",
   "DRAFT_RUN_ID",
   "DRAFT_HASH",
   "HR_APPROVED_BY",
   "HR_APPROVED_AT",
   "ACCOUNTS_APPROVED_BY",
   "ACCOUNTS_APPROVED_AT",
   "LOCKED_AT",
   "LOCK_ID"
  ],
  "FEED_STATUS": [
   "PERIOD",
   "FEED",
   "STATUS",
   "MARKED_BY",
   "MARKED_AT",
   "NOTE"
  ],
  "PAYROLL_CATEGORY_CONFIG": [
   "CATEGORY_CODE",
   "DISPLAY_NAME",
   "CALC_METHOD",
   "SITE",
   "PAYSLIP",
   "PAYSLIP_TEMPLATE_KEY",
   "RATE_SOURCE",
   "ACTIVE",
   "APPROVED_BY",
   "APPROVED_AT"
  ],
  "STATUTORY_CONFIG": [
   "KEY",
   "VALUE",
   "NOTE",
   "EFFECTIVE_FROM",
   "EFFECTIVE_TO",
   "VERSION",
   "APPROVED_BY",
   "APPROVED_AT"
  ],
  "EFFICIENCY_CONFIG": [
   "EFFICIENCY_PERCENT_EXACT",
   "INCENTIVE_SLAB_INR",
   "BASIS",
   "SOURCE",
   "IMPLEMENTATION_STATE",
   "NOTE"
  ],
  "PT_EXEMPTIONS": [
   "EMP_ID",
   "REASON",
   "EFFECTIVE_FROM",
   "EFFECTIVE_TO",
   "APPROVED_BY"
  ],
  "HOLIDAY_CALENDAR": [
   "DATE",
   "SITE",
   "HOLIDAY_NAME",
   "PAID"
  ],
  "EMPLOYEE_MASTER": [
   "EMP_ID",
   "EMPLOYEE_NAME",
   "EMAIL_ID",
   "DOJ_AS_SOURCE",
   "PAYROLL_CATEGORY",
   "STATUS_AS_SOURCE",
   "DEPARTMENT",
   "DESIGNATION",
   "PLANT_TO_VERIFY",
   "MANAGER_EMAIL_TO_VERIFY",
   "STATUTORY_PROFILE_TO_VERIFY",
   "SOURCE_RECORD_KEY",
   "SOURCE_TAB",
   "SOURCE_ROW",
   "DUPLICATE_FLAG",
   "VALIDATION_STATE",
   "SOURCE_SNAPSHOT_DATE",
   "HR_SIGNOFF_BY",
   "HR_SIGNOFF_AT",
   "REVIEW_NOTE",
   "LAST_WORKING_DAY"
  ],
  "SALARY_STRUCTURE": [
   "EMP_ID",
   "PAYROLL_CATEGORY",
   "SOURCE_PAYROLL_MONTH",
   "EFFECTIVE_FROM",
   "EFFECTIVE_TO",
   "EMPLOYMENT_STATUS_AT_SOURCE",
   "BASIC_PM_INR",
   "HRA_PM_INR",
   "CONVEYANCE_PM_INR",
   "EDUCATION_PM_INR",
   "MEDICAL_PM_INR",
   "PRO_DEV_PM_INR",
   "COMMUNICATION_PM_INR",
   "UNIFORM_PM_INR",
   "WASHING_PM_INR",
   "HEAT_MASTER_INR",
   "VDA_MASTER_INR",
   "PRODUCTION_MASTER_INR",
   "FIXED_GROSS_PM_AS_SOURCE_INR",
   "CTC_PA_AS_SOURCE_INR",
   "CTC_PM_AS_SOURCE_INR",
   "SOURCE_TAB",
   "SOURCE_ROW",
   "SOURCE_ROW_KEY",
   "VERSION_STATE",
   "HR_APPROVED_BY",
   "HR_APPROVED_AT",
   "VALIDATION_NOTE"
  ],
  "PAYROLL_RATE_PROFILE": [
   "EMP_ID",
   "PAYROLL_CATEGORY",
   "PAY_BASIS",
   "RATE_AMOUNT_INR",
   "MONTHLY_GROSS_INR",
   "ATTENDANCE_REQUIRED",
   "WORKING_DAYS_REQUIRED",
   "PRESENT_DAYS_REQUIRED",
   "WORKED_DAYS_REQUIRED",
   "OT_METHOD",
   "BASELINE_MONTH",
   "SOURCE_MONTH",
   "SOURCE_USAGE",
   "SOURCE_SPREADSHEET_ID",
   "SOURCE_SHEET",
   "SOURCE_ROW",
   "VERSION_STATE",
   "NOTE",
   "EFFECTIVE_FROM",
   "EFFECTIVE_TO",
   "HR_APPROVED_BY",
   "HR_APPROVED_AT"
  ],
  "EMPLOYEE_STATUTORY_IDS": [
   "EMP_ID",
   "UAN",
   "ESI_NO",
   "PAN",
   "BANK_NAME",
   "BANK_ACCOUNT",
   "IFSC"
  ],
  "INPUT_OT": [
   "PAYROLL_MONTH",
   "EMP_ID",
   "OT_HOURS",
   "SOURCE_REF",
   "APPROVAL_STATUS",
   "ENTERED_AT",
   "SOURCE_CASE_NOS",
   "SOURCE_EVENT_COUNT",
   "DATE_RANGE",
   "OT_KEY",
   "OT_DATE",
   "SOURCE_ROW",
   "NORMALIZER_VERSION",
   "ELIGIBILITY",
   "EXCEPTION_REASON"
  ],
  "INPUT_CANTEEN": [
   "PAYROLL_MONTH",
   "EMP_ID",
   "AMOUNT_INR",
   "SOURCE",
   "SOURCE_REF",
   "KEY",
   "STATUS",
   "ENTERED_AT",
   "REMARKS"
  ],
  "INPUT_EFFICIENCY": [
   "PAYROLL_MONTH",
   "EMP_ID",
   "EFFICIENCY_PCT",
   "PHYSICAL_PRESENT_DAYS_OVERRIDE",
   "SOURCE",
   "SOURCE_REF",
   "KEY",
   "STATUS",
   "ENTERED_AT",
   "REMARKS"
  ],
  "INPUT_ADVANCE": [
   "PAYROLL_MONTH",
   "EMP_ID",
   "EMPLOYEE_NAME_DISPLAY",
   "ADVANCE_TYPE",
   "ADVANCE_DATE",
   "ORIGINAL_ADVANCE_INR",
   "OPENING_BALANCE_INR",
   "RECOVERY_THIS_MONTH_INR",
   "CLOSING_BALANCE_INR",
   "ACCOUNTS_LEDGER_REFERENCE",
   "SOURCE_BATCH_ID",
   "APPROVAL_STATUS",
   "APPROVED_BY",
   "ENTERED_AT",
   "REMARKS"
  ],
  "INPUT_SOCIETY": [
   "PAYROLL_MONTH",
   "EMP_ID",
   "EMPLOYEE_NAME_DISPLAY",
   "SOCIETY_NAME",
   "LOAN_REFERENCE",
   "GENERAL_EMI_INR",
   "EMERGENCY_EMI_INR",
   "EDUCATION_EMI_INR",
   "SHARES_OTHER_INR",
   "TOTAL_RECOVERY_INR",
   "OUTSTANDING_BALANCE_INR",
   "APPROVAL_STATUS",
   "SOURCE_BATCH_ID",
   "REMARKS"
  ],
  "INPUT_ADJUSTMENTS": [
   "ENTRY_ID",
   "PAYROLL_MONTH",
   "EMP_ID",
   "EMPLOYEE_NAME_DISPLAY",
   "ADJUSTMENT_TYPE",
   "SIGNED_AMOUNT_INR",
   "REASON",
   "SOURCE_REFERENCE",
   "APPROVAL_STATUS",
   "APPROVED_BY",
   "APPROVED_AT",
   "REVERSAL_OF_ENTRY_ID",
   "ENTERED_BY",
   "ENTERED_AT"
  ],
  "INPUT_LEAVE": [
   "PAYROLL_MONTH",
   "EMP_ID",
   "LEAVE_TYPE",
   "DAYS",
   "FROM_DATE",
   "TO_DATE",
   "SOURCE_REF",
   "CASE_NO",
   "KEY",
   "STATUS",
   "EXCEPTION_REASON",
   "NORMALIZER_VERSION",
   "ENTERED_AT"
  ],
  "ATTENDANCE_DAILY": [
   "PERIOD",
   "DATE",
   "SITE",
   "EMP_ID",
   "CODE",
   "SOURCE",
   "SOURCE_REF",
   "KEY",
   "STATUS",
   "REJECT_REASON",
   "ENTERED_AT"
  ],
  "INPUT_ATTENDANCE": [
   "PAYROLL_MONTH",
   "EMP_ID",
   "PAYROLL_CATEGORY",
   "WORKING_DAYS",
   "PRESENT_DAYS",
   "WEEK_OFF",
   "PH",
   "EL_AVAILED",
   "CL_AVAILED",
   "SL_AVAILED",
   "PAID_LEAVE_OTHER",
   "WORKED_DAYS",
   "PAYABLE_DAYS",
   "APPROVAL_STATUS",
   "APPROVED_BY",
   "SOURCE_REF",
   "ENTERED_AT",
   "REMARKS",
   "PHYSICAL_PRESENT_DAYS",
   "ABSENT_LWP_DAYS",
   "GENERATED_VALUES_JSON",
   "HR_OVERRIDE",
   "OVERRIDE_REASON",
   "ROW_KEY",
   "REGISTER_DAYS_PRESENT",
   "REGISTER_INCLUDES_WO",
   "ENTERED_BY"
  ],
  "ATTENDANCE_COMPARISON": [
   "PERIOD",
   "EMP_ID",
   "NAME",
   "POPULATION",
   "DAILY_PRESENT",
   "REGISTER_PRESENT",
   "DIFF",
   "STATUS",
   "HR_DECIDED_DAYS",
   "HR_REASON",
   "HR_BY",
   "HR_AT",
   "OWNER_DECISION",
   "OWNER_BY",
   "OWNER_AT",
   "HR_STAMPED_DAYS"
  ],
  "PAYROLL_READINESS": [
   "PERIOD",
   "POPULATION",
   "CHECK",
   "STATUS",
   "DETAIL",
   "CHECKED_AT"
  ],
  "PAYROLL_DRAFT": [
   "RUN_ID",
   "PERIOD",
   "POPULATION",
   "EMP_ID",
   "EMPLOYEE_NAME",
   "DEPARTMENT",
   "DESIGNATION",
   "WORKING_DAYS",
   "PRESENT_DAYS",
   "PHYSICAL_PRESENT_DAYS",
   "WO_DAYS",
   "PH_DAYS",
   "EL",
   "CL",
   "SL",
   "PAID_LEAVE_OTHER",
   "ABSENT_LWP_DAYS",
   "WORKED_PAYABLE_DAYS",
   "FIXED_GROSS",
   "PAY_BASIS",
   "RATE",
   "BASIC",
   "HRA",
   "CONVEYANCE",
   "EDUCATION",
   "MEDICAL",
   "PRO_DEV",
   "COMMUNICATION",
   "UNIFORM",
   "WASHING",
   "HEAT",
   "VDA",
   "PRODUCTION_ALLOWANCE",
   "GROSS_EARNINGS",
   "OT_HOURS",
   "OT_AMOUNT",
   "ARREARS",
   "DISPATCH_INCENTIVE",
   "OTHER_ALLOWANCE",
   "LEAVE_ENCASHMENT",
   "PRODUCTION_INCENTIVE",
   "OT_EXTRA_WORK",
   "TOTAL_EARNINGS",
   "PF_EMPLOYEE",
   "ESI_EMPLOYEE",
   "PT",
   "MLWF",
   "CANTEEN",
   "SOCIETY",
   "ADVANCE",
   "TDS",
   "EFFICIENCY_PCT",
   "EFFICIENCY_ELIGIBLE_AMOUNT",
   "EFFICIENCY_DEDUCTION",
   "OTHER_DEDUCTION",
   "TOTAL_DEDUCTIONS",
   "NET_PAY",
   "EMPLOYER_PF",
   "EMPLOYER_ESI",
   "BONUS_PROVISION",
   "GRATUITY_PROVISION",
   "FLAGS",
   "CALC_VERSION",
   "CALCULATED_AT"
  ],
  "PAYROLL_EXCEPTIONS": [
   "RUN_ID",
   "PERIOD",
   "POPULATION",
   "EMP_ID",
   "SEVERITY",
   "CODE",
   "MESSAGE"
  ],
  "PAYROLL_RECON": [
   "PERIOD",
   "POPULATION",
   "HEADCOUNT",
   "TOTAL_GROSS",
   "TOTAL_DEDUCTIONS",
   "TOTAL_NET",
   "PREV_PERIOD_NET",
   "DELTA_PCT",
   "RUN_ID"
  ],
  "PAYROLL_SUPPLEMENTARY": [
   "PERIOD",
   "POPULATION",
   "SUPP_ID",
   "EMP_IDS",
   "HASH",
   "STATUS",
   "CREATED_BY",
   "CREATED_AT",
   "HR_APPROVED_BY",
   "HR_APPROVED_AT",
   "ACCOUNTS_APPROVED_BY",
   "ACCOUNTS_APPROVED_AT",
   "LOCK_ID",
   "LOCKED_AT"
  ],
  "PAYROLL_LOCKED": [
   "LOCK_ID",
   "RUN_ID",
   "PERIOD",
   "POPULATION",
   "EMP_ID",
   "EMPLOYEE_NAME",
   "DEPARTMENT",
   "DESIGNATION",
   "WORKING_DAYS",
   "PRESENT_DAYS",
   "PHYSICAL_PRESENT_DAYS",
   "WO_DAYS",
   "PH_DAYS",
   "EL",
   "CL",
   "SL",
   "PAID_LEAVE_OTHER",
   "ABSENT_LWP_DAYS",
   "WORKED_PAYABLE_DAYS",
   "FIXED_GROSS",
   "PAY_BASIS",
   "RATE",
   "BASIC",
   "HRA",
   "CONVEYANCE",
   "EDUCATION",
   "MEDICAL",
   "PRO_DEV",
   "COMMUNICATION",
   "UNIFORM",
   "WASHING",
   "HEAT",
   "VDA",
   "PRODUCTION_ALLOWANCE",
   "GROSS_EARNINGS",
   "OT_HOURS",
   "OT_AMOUNT",
   "ARREARS",
   "DISPATCH_INCENTIVE",
   "OTHER_ALLOWANCE",
   "LEAVE_ENCASHMENT",
   "PRODUCTION_INCENTIVE",
   "OT_EXTRA_WORK",
   "TOTAL_EARNINGS",
   "PF_EMPLOYEE",
   "ESI_EMPLOYEE",
   "PT",
   "MLWF",
   "CANTEEN",
   "SOCIETY",
   "ADVANCE",
   "TDS",
   "EFFICIENCY_PCT",
   "EFFICIENCY_ELIGIBLE_AMOUNT",
   "EFFICIENCY_DEDUCTION",
   "OTHER_DEDUCTION",
   "TOTAL_DEDUCTIONS",
   "NET_PAY",
   "EMPLOYER_PF",
   "EMPLOYER_ESI",
   "BONUS_PROVISION",
   "GRATUITY_PROVISION",
   "FLAGS",
   "CALC_VERSION",
   "CALCULATED_AT"
  ],
  "PAYSLIP_REGISTER": [
   "LOCK_ID",
   "PERIOD",
   "EMP_ID",
   "POPULATION",
   "DOC_ID",
   "PDF_ID",
   "PDF_URL",
   "GENERATED_AT",
   "STATUS"
  ],
  "PAYSLIP_EMAIL_LOG": [
   "LOCK_ID",
   "PERIOD",
   "EMP_ID",
   "TO_EMAIL",
   "PDF_ID",
   "STATUS",
   "ATTEMPTED_AT",
   "ERROR"
  ],
  "AUDIT_LOG": [
   "Timestamp",
   "Module",
   "Status",
   "User",
   "Message"
  ]
 },
 "control": [
  [
   "HR_APPROVER_EMAIL",
   "hr@varshaforgings.com",
   "HR approver (state machine)"
  ],
  [
   "ACCOUNTS_APPROVER_EMAIL",
   "accounts@varshaforgings.com",
   "Accounts approver (state machine)"
  ],
  [
   "MIN_PERIOD",
   "2026-09",
   "Earliest period any write path accepts"
  ],
  [
   "PAYSLIP_TEMPLATE_STAFF_ID",
   "1T7kwVNmOczXk4_-4OstofWOPSWGuEQxg7hKFo-Jci_w",
   "Staff payslip template Doc"
  ],
  [
   "PAYSLIP_TEMPLATE_WORKER_ID",
   "1MSmi8qVRL8SI8-svVihasYbLFGNo8Xkzko4VUaIX8SU",
   "Worker payslip template Doc"
  ],
  [
   "PAYSLIP_FOLDER_ID",
   "",
   "Blank = payslip step blocked"
  ],
  [
   "EMAIL_RELEASE_ENABLED",
   "FALSE",
   "Payslip email release switch"
  ],
  [
   "VFL_WEEKLY_OFF",
   "SUN",
   "Weekly off used to default blank attendance"
  ],
  [
   "PUNE_WEEKLY_OFF",
   "SUN",
   "Weekly off used to default blank attendance"
  ],
  [
   "OT_SOURCE_SPREADSHEET_ID",
   "",
   "Blank = read the OT form responses from a local tab of this spreadsheet; set only to read an external response spreadsheet"
  ],
  [
   "OT_SOURCE_TAB",
   "OT_FORM_RESPONSES",
   "OT form-response tab (local; falls back to Overtime_Form if absent). With an external ID: the tab there (default Form Responses 1)"
  ],
  [
   "OT_WINDOW_START_2026-09",
   "2026-08-26",
   "one-time catch-up: Aug salary paid OT to 25-Aug"
  ],
  [
   "LEAVE_WINDOW_START_2026-09",
   "2026-08-26",
   "one-time catch-up: Aug payroll counted leave to 25-Aug; default leave window = calendar month"
  ],
  [
   "OWNER_APPROVER_EMAIL",
   "yash.munot@gmail.com",
   "confirm owner email (owner approval of attendance disputes)"
  ],
  [
   "REGISTER_ENTRY_EMAILS",
   "",
   "Extra people (comma separated) who may submit the monthly attendance register; HR_APPROVER_EMAIL and OWNER_APPROVER_EMAIL always may"
  ],
  [
   "LEAVE_SOURCE_SPREADSHEET_ID",
   "1pwVE0XKqAhAKHbyqtlF9GzfuGnidnZuw2zKbtMjUz9Q",
   "Leave application spreadsheet (read-only; give the script runner view access). Blank = read a local tab of this spreadsheet"
  ],
  [
   "LEAVE_SOURCE_TAB",
   "Leave_Applications",
   "Leave form-response tab in the leave spreadsheet (or the local tab when the ID is blank)"
  ]
 ],
 "statutoryDefaults": [
  [
   "PT_FEB_AMOUNT",
   "300",
   "February PT amount"
  ],
  [
   "MLWF_MONTHS",
   "6,12",
   "Months MLWF is deducted"
  ],
  [
   "STAFF_OT_MULTIPLIER",
   "2",
   ""
  ],
  [
   "WORKER_OT_MULTIPLIER",
   "2",
   ""
  ],
  [
   "STAFF_PF_WAGE_COMPONENTS",
   "BASIC,CONVEYANCE,EDUCATION,MEDICAL",
   ""
  ],
  [
   "STAFF_COMPONENT_PCTS",
   "{\"BASIC\":0.40,\"HRA\":0.24,\"CONVEYANCE\":0.06,\"MEDICAL\":0.06,\"EDUCATION\":0.06,\"PRO_DEV\":0.03,\"COMMUNICATION\":0.02,\"UNIFORM\":0.04,\"WASHING\":0.09}",
   ""
  ],
  [
   "EMPLOYER_PF_RATE_STAFF",
   "0.1301",
   ""
  ],
  [
   "EMPLOYER_PF_RATE_WORKER",
   "0.1301",
   ""
  ],
  [
   "BONUS_RATE_STAFF",
   "0.0833",
   ""
  ],
  [
   "GRATUITY_RATE_STAFF",
   "0.0483",
   ""
  ],
  [
   "BONUS_RATE_WORKER",
   "0.18",
   ""
  ],
  [
   "GRATUITY_RATE_WORKER",
   "0.0481",
   ""
  ]
 ],
 "categories": [
  {
   "CATEGORY_CODE": "STAFF",
   "DISPLAY_NAME": "Staff",
   "CALC_METHOD": "STAFF",
   "SITE": "VFL",
   "PAYSLIP": "Y",
   "PAYSLIP_TEMPLATE_KEY": "STAFF",
   "RATE_SOURCE": "SALARY_STRUCTURE",
   "ACTIVE": "Y"
  },
  {
   "CATEGORY_CODE": "PERMANENT_WORKER",
   "DISPLAY_NAME": "Permanent worker",
   "CALC_METHOD": "PERMANENT_WORKER",
   "SITE": "VFL",
   "PAYSLIP": "Y",
   "PAYSLIP_TEMPLATE_KEY": "WORKER",
   "RATE_SOURCE": "SALARY_STRUCTURE",
   "ACTIVE": "Y"
  },
  {
   "CATEGORY_CODE": "CONSULTANT",
   "DISPLAY_NAME": "Consultant",
   "CALC_METHOD": "CONSULTANT",
   "SITE": "VFL",
   "PAYSLIP": "N",
   "PAYSLIP_TEMPLATE_KEY": "",
   "RATE_SOURCE": "RATE_PROFILE",
   "ACTIVE": "Y"
  },
  {
   "CATEGORY_CODE": "PUNE_STAFF",
   "DISPLAY_NAME": "Pune staff",
   "CALC_METHOD": "PUNE_STAFF",
   "SITE": "PUNE",
   "PAYSLIP": "N",
   "PAYSLIP_TEMPLATE_KEY": "",
   "RATE_SOURCE": "RATE_PROFILE",
   "ACTIVE": "Y"
  }
 ],
 "ptExemptions": [
  "VFL4021",
  "VFL4014",
  "VFL4063"
 ]
}""")
# SCHEMA-END

MIN_PERIOD = "2026-09"
TEXT_COLS = {"EMP_ID", "PAYROLL_MONTH", "PERIOD", "EFFECTIVE_FROM", "EFFECTIVE_TO", "DOJ_AS_SOURCE", "LAST_WORKING_DAY", "UAN",
             "ESI_NO", "PAN", "BANK_NAME", "BANK_ACCOUNT", "IFSC", "VALUE", "KEY", "SOURCE_ROW", "SOURCE_SNAPSHOT_DATE"}
FORBIDDEN_HEADER = re.compile(r"password|aadhaa?r|mobile|phone", re.I)
INPUT_TABS = ["INPUT_ATTENDANCE", "INPUT_OT", "INPUT_CANTEEN", "INPUT_EFFICIENCY", "INPUT_ADVANCE", "INPUT_SOCIETY",
              "INPUT_ADJUSTMENTS", "INPUT_LEAVE"]
OUTPUT_ORDER = ["PAYROLL_CONTROL", "PAYROLL_CATEGORY_CONFIG", "PAYROLL_PERIOD_CATEGORY", "STATUTORY_CONFIG", "EFFICIENCY_CONFIG",
                "PT_EXEMPTIONS", "HOLIDAY_CALENDAR", "EMPLOYEE_MASTER", "SALARY_STRUCTURE", "PAYROLL_RATE_PROFILE",
                "EMPLOYEE_STATUTORY_IDS"] + INPUT_TABS


def clean(v):
    """Cell value -> plain value: integral floats become ints, datetimes ISO text, blanks ''."""
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v) and abs(v) < 1e15:
        return int(v)
    if hasattr(v, "isoformat"):
        return v.isoformat(sep=" ")[:19] if hasattr(v, "hour") else v.isoformat()
    return v


def as_text(v):
    v = clean(v)
    return "" if v == "" else str(v).strip()


def table(ws, header_row=1, first_row=None):
    """Rows of a tab as dicts keyed by the header row (blank rows skipped)."""
    first_row = first_row or header_row + 1
    hdr = [as_text(c.value) for c in ws[header_row]]
    out = []
    for row in ws.iter_rows(min_row=first_row, values_only=True):
        if all(x in (None, "") for x in row):
            continue
        out.append({h: clean(row[i]) for i, h in enumerate(hdr) if h and i < len(row)})
    return out


def norm(s):
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())


def combined_header(ws, rows):
    """Header names of a split header (the first non-blank of the given rows per column)."""
    names = []
    for i in range(1, ws.max_column + 1):
        v = ""
        for r in rows:
            x = ws.cell(r, i).value
            if x not in (None, ""):
                v = norm(x)
                break
        names.append(v)
    return names


IDENT_FIELDS = {"UAN": ["uan"], "ESI_NO": ["esino"], "PAN": ["pan", "panno"], "BANK_NAME": ["bankname"],
                "BANK_ACCOUNT": ["accountno", "accountnumber", "bankaccountno"], "IFSC": ["ifsc"]}


def read_identity(ws, header_rows, first_row):
    """{EMP_ID: {field: text}} from a hidden master: only the EMP CODE column and the 6 identity columns are read."""
    names = combined_header(ws, header_rows)

    def col(cands):
        for c in cands:
            if c in names:
                return names.index(c) + 1
        return None

    emp = col(["empcode", "empid", "employeecode"])
    cols = {f: col(c) for f, c in IDENT_FIELDS.items()}
    out = {}
    if not emp:
        return out
    for r in range(first_row, ws.max_row + 1):
        eid = as_text(ws.cell(r, emp).value)
        if not eid:
            continue
        rec = out.setdefault(eid, {})
        for f, c in cols.items():
            v = as_text(ws.cell(r, c).value) if c else ""
            if v and v.upper() not in ("NA", "N/A", "#N/A", "NONE", "-") and not v.startswith("#") and not rec.get(f):
                rec[f] = v
    return out


def write_tab(wb, name, headers, rows):
    ws = wb.create_sheet(name)
    ws.append(headers)
    text_idx = [i + 1 for i, h in enumerate(headers) if h in TEXT_COLS]
    for r in rows:
        ws.append([clean(r.get(h, "")) if not isinstance(r.get(h, ""), (list, dict)) else "" for h in headers])
    for i in text_idx:
        for row in range(1, ws.max_row + 1):
            cell = ws.cell(row, i)
            if row > 1 and cell.value not in (None, ""):
                cell.value = str(cell.value)
                cell.data_type = "s"
            cell.number_format = "@"
    ws.freeze_panes = "A2"
    return ws


def main(argv):
    if len(argv) != 3:
        sys.exit(__doc__)
    old = openpyxl.load_workbook(argv[1], data_only=True)
    H = SCHEMA["headers"]
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    counts = {}

    def put(name, rows):
        write_tab(wb, name, H[name], rows)
        counts[name] = len(rows)

    put("PAYROLL_CONTROL", [{"KEY": k, "VALUE": v, "NOTE": n, "UPDATED_AT": ""} for k, v, n in SCHEMA["control"]])
    put("PAYROLL_CATEGORY_CONFIG", [dict(c, APPROVED_BY="", APPROVED_AT="") for c in SCHEMA["categories"]])
    put("PAYROLL_PERIOD_CATEGORY", [{"PAYROLL_MONTH": MIN_PERIOD, "PAYROLL_CATEGORY": c["CATEGORY_CODE"], "WORKING_DAYS": "",
                                     "STATUS": "PENDING", "NOTE": "HR enters WORKING_DAYS for %s." % MIN_PERIOD}
                                    for c in SCHEMA["categories"]])
    stat = [dict(r, EFFECTIVE_FROM=MIN_PERIOD, VERSION=1, APPROVED_BY="", APPROVED_AT="")
            for r in table(old["STATUTORY_CONFIG"])]
    have = {r["KEY"] for r in stat}
    for k, v, n in SCHEMA["statutoryDefaults"]:
        if k not in have:
            stat.append({"KEY": k, "VALUE": (float(v) if re.fullmatch(r"-?\d+(\.\d+)?", v) else v), "NOTE": n,
                         "EFFECTIVE_FROM": MIN_PERIOD, "VERSION": 1, "APPROVED_BY": "", "APPROVED_AT": ""})
    put("STATUTORY_CONFIG", stat)
    put("EFFICIENCY_CONFIG", table(old["EFFICIENCY_CONFIG"]))
    put("PT_EXEMPTIONS", [{"EMP_ID": i, "REASON": "carried from Aug worker template R17", "EFFECTIVE_FROM": MIN_PERIOD,
                           "EFFECTIVE_TO": "", "APPROVED_BY": "SETUP_SEED"} for i in SCHEMA["ptExemptions"]])
    put("HOLIDAY_CALENDAR", [])
    master = table(old["EMPLOYEE_MASTER"])
    put("EMPLOYEE_MASTER", [dict(r, LAST_WORKING_DAY="") for r in master])
    put("SALARY_STRUCTURE", [dict(r, HR_APPROVED_BY="", HR_APPROVED_AT="") for r in table(old["SALARY_STRUCTURE"])])
    put("PAYROLL_RATE_PROFILE", [dict(r, EFFECTIVE_FROM="2026-08-01", EFFECTIVE_TO="", HR_APPROVED_BY="", HR_APPROVED_AT="")
                                 for r in table(old["PAYROLL_RATE_PROFILE"])])

    # statutory IDs: the hidden RAW masters (identity headers on row 2, EMP CODE on row 4, data from row 5); the
    # *_VALUES copies and the *_AUG26 tabs (flat header on row 1, data from row 4) only fill fields that are still blank
    ident = {}
    sources = [("RAW_STAFF_MASTER", (4, 2), 5), ("RAW_WORKER_MASTER", (4, 2), 5), ("RAW_STAFF_MASTER_VALUES", (4, 2), 5),
               ("RAW_WORKER_MASTER_VALUES", (4, 2), 5), ("RAW_STAFF_AUG26", (1,), 4), ("RAW_WORKER_AUG26", (1,), 4)]
    used = []
    for name, hrows, first in sources:
        if name not in old.sheetnames:
            continue
        got = read_identity(old[name], hrows, first)
        used.append(name)
        for eid, rec in got.items():
            tgt = ident.setdefault(eid, {})
            for f, v in rec.items():
                if v and not tgt.get(f):
                    tgt[f] = v
    known = {as_text(r.get("EMP_ID")) for r in master}
    id_rows = [dict(EMP_ID=e, **{f: rec.get(f, "") for f in IDENT_FIELDS}) for e, rec in sorted(ident.items())
               if e in known and any(rec.values())]
    put("EMPLOYEE_STATUTORY_IDS", id_rows)
    for t in INPUT_TABS:
        put(t, [])

    # ---------------------------------------------------------------- self-check (counts only)
    for ws in wb.worksheets:
        for c in ws[1]:
            assert not FORBIDDEN_HEADER.search(str(c.value or "")), "forbidden column %s in %s" % (c.value, ws.title)
    assert [ws.title for ws in wb.worksheets] == OUTPUT_ORDER, "unexpected tab list"
    assert H["EMPLOYEE_STATUTORY_IDS"] == ["EMP_ID", "UAN", "ESI_NO", "PAN", "BANK_NAME", "BANK_ACCOUNT", "IFSC"]
    active = [as_text(r.get("EMP_ID")) for r in master if as_text(r.get("STATUS_AS_SOURCE")).lower() == "active"]
    assert len(active) == len(set(active)), "duplicate EMP_ID among active employees"
    assert len({r["EMP_ID"] for r in id_rows}) == len(id_rows), "duplicate EMP_ID in EMPLOYEE_STATUTORY_IDS"
    assert counts["EMPLOYEE_MASTER"] == len(master) and counts["PT_EXEMPTIONS"] == 3
    active_set = set(active)
    covered = len(active_set & {r["EMP_ID"] for r in id_rows})
    wb.save(argv[2])
    for name in OUTPUT_ORDER:
        print("%-26s %d rows" % (name, counts[name]))
    print("active employees %d, with statutory IDs %d (sources read: %s)" % (len(active), covered, ", ".join(used)))
    print("self-check OK")


if __name__ == "__main__":
    main(sys.argv)
